"""
DOCX and text parser for quiz questions.
Follows the standard:
- Question blocks delimited by '+++'
- Within each block, parts delimited by '==='
- Part 0: Question text
- Parts 1..N: Answer options (min 2, max 10)
- Exactly one correct answer marked with '#', '+', '✅', or '✓'
"""

import re
from pathlib import Path
from typing import BinaryIO, Optional, Union
from docx import Document

MARKERS = ["#", "✅", "✓", "+"]


def parse_question_block(block: str) -> Optional[dict]:
    """
    Parse a single question block delimited by '==='.
    Returns a dict with 'question', 'options', and 'correct_option_id' or None if malformed.
    """
    raw_parts = block.split("===")
    parts = [p.strip() for p in raw_parts if p.strip()]
    if len(parts) < 3:
        return None

    question_text = parts[0]
    raw_options = parts[1:]

    # Telegram poll limits: 2 to 10 options
    if len(raw_options) < 2 or len(raw_options) > 10:
        return None

    # Step 1: Check prefix markers (e.g. #Option, + Option, ✅ Option, ✓ Option)
    prefix_marked = []
    for idx, opt in enumerate(raw_options):
        s = opt.strip()
        for m in MARKERS:
            if s.startswith(m):
                prefix_marked.append((idx, m, s[len(m):].strip()))
                break

    if len(prefix_marked) == 1:
        correct_idx, marker, cleaned = prefix_marked[0]
        cleaned_options = []
        for i, opt in enumerate(raw_options):
            if i == correct_idx:
                cleaned_options.append(cleaned)
            else:
                cleaned_options.append(opt.strip())

        if not cleaned or any(not opt for opt in cleaned_options):
            return None

        # Truncate to Telegram limits if necessary
        if len(question_text) > 300:
            question_text = question_text[:297] + "..."
        cleaned_options = [opt[:100] for opt in cleaned_options]

        return {
            "question": question_text,
            "options": cleaned_options,
            "correct_option_id": correct_idx,
        }

    # Step 2: Check marker containment anywhere if prefix marker not unique
    contain_marked = []
    for idx, opt in enumerate(raw_options):
        s = opt.strip()
        for m in MARKERS:
            if m in s:
                cleaned = s.replace(m, "", 1).strip()
                contain_marked.append((idx, m, cleaned))
                break

    if len(contain_marked) == 1:
        correct_idx, marker, cleaned = contain_marked[0]
        cleaned_options = []
        for i, opt in enumerate(raw_options):
            if i == correct_idx:
                cleaned_options.append(cleaned)
            else:
                cleaned_options.append(opt.strip())

        if not cleaned or any(not opt for opt in cleaned_options):
            return None

        # Truncate to Telegram limits if necessary
        if len(question_text) > 300:
            question_text = question_text[:297] + "..."
        cleaned_options = [opt[:100] for opt in cleaned_options]

        return {
            "question": question_text,
            "options": cleaned_options,
            "correct_option_id": correct_idx,
        }

    # 0 or >1 correct markers -> ignore malformed block gracefully
    return None


def parse_text(text: str) -> list[dict]:
    """
    Parse questions from raw text containing blocks separated by '+++'.
    """
    blocks = text.split("+++")
    questions = []
    for block in blocks:
        if not block.strip():
            continue
        q = parse_question_block(block)
        if q:
            questions.append(q)
    return questions


def extract_text_from_docx(file_source: Union[str, Path, BinaryIO]) -> str:
    """
    Extract text content from a docx file path or binary stream.
    """
    doc = Document(file_source)
    lines = []
    for p in doc.paragraphs:
        lines.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                lines.append(" === ".join(row_texts))
    return "\n".join(lines)


def parse_docx(file_source: Union[str, Path, BinaryIO]) -> list[dict]:
    """
    Read and parse questions from a .docx file or stream.
    """
    text = extract_text_from_docx(file_source)
    return parse_text(text)
